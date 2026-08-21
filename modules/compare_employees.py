# -*- coding: utf-8 -*-
"""Сверка HR Excel и e-SEN + загрузка HR Excel с автоопределением колонок."""
import json
import os
import re
from datetime import datetime
from pathlib import Path

from tkinter import filedialog, messagebox

from modules.name_match import names_match
from modules.export_reports import export_missing_separate_files_by_departments

PROJECT_ROOT = Path(__file__).resolve().parent.parent
HR_FILE = PROJECT_ROOT / "database" / "hr_employees.json"
ESEN_FILE = PROJECT_ROOT / "database" / "esen_employees.json"
RESULT_FILE = PROJECT_ROOT / "database" / "compare_result.json"
HISTORY_FILE = PROJECT_ROOT / "database" / "sync_history.json"

DEP_KEYS = ("отдел", "подраздел", "департамент", "служба", "сектор", "цех", "управлен", "бөлім", "department")


# ================= СЛУЖЕБНЫЕ =================
def load_json(path):
    if not os.path.exists(path):
        print(f"Файл не найден: {path}")
        return []
    with open(path, "r", encoding="utf-8-sig") as f:
        return json.load(f)


def save_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)


def append_sync_history(result):
    history = []
    if os.path.exists(HISTORY_FILE):
        try:
            with open(HISTORY_FILE, "r", encoding="utf-8-sig") as f:
                history = json.load(f)
        except Exception:
            history = []
    if not isinstance(history, list):
        history = []
    history.append({
        "synced_at": result.get("synced_at", ""),
        "total_hr": result.get("total_hr", 0),
        "total_esen": result.get("total_esen", 0),
        "matched": len(result.get("matched", [])),
        "only_hr": len(result.get("only_hr", [])),
        "only_esen": len(result.get("only_esen", [])),
    })
    history = history[-100:]
    save_json(HISTORY_FILE, history)


# ================= ЗАГРУЗКА HR EXCEL =================
def norm(s):
    return str(s if s is not None else "").strip().lower().replace("ё", "е")


def _read_rows(path):
    path = str(path)
    if path.lower().endswith(".csv"):
        import csv
        with open(path, "r", encoding="utf-8-sig", errors="ignore") as f:
            return [row for row in csv.reader(f)]
    try:
        import pandas as pd
        return pd.read_excel(path, header=None, dtype=str).fillna("").values.tolist()
    except ImportError:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows


def _col_values(rows, j):
    out = []
    for r in rows:
        v = str(r[j]).strip() if j < len(r) and r[j] is not None else ""
        if v:
            out.append(v)
    return out


def _stats(rows, j):
    vals = _col_values(rows, j)
    if len(vals) < 2:
        return None
    digits = sum(1 for v in vals if re.fullmatch(r"[\d\-\+\. ]+", v))
    caps = sum(1 for v in vals if sum(1 for w in v.split() if w[:1].isupper()) >= 2)
    return {
        "digit": digits / len(vals),
        "words": sum(len(v.split()) for v in vals) / len(vals),
        "caps": caps / len(vals),
        "card": len(set(v.lower() for v in vals)),
        "dep_kw": sum(1 for v in vals if any(k in v.lower() for k in DEP_KEYS)),
    }


def _is_fio_header(c):
    return ("фио" in c or "сотрудник" in c or "фамилия" in c or "ф.и.о" in c
            or "тегі" in c or "аты" in c or "қызметкер" in c or "full name" in c)


def _is_pos_header(c):
    return ("должност" in c or "лауазым" in c or "позици" in c or "position" in c)


def _is_dep_header(c):
    return any(k in c for k in DEP_KEYS)


def detect_columns(rows):
    for i, row in enumerate(rows[:10]):
        cells = [norm(c) for c in row]
        if any(_is_fio_header(c) for c in cells):
            idx = {"fio": None, "pos": None, "dep": None}
            for j, c in enumerate(cells):
                if not c:
                    continue
                if idx["fio"] is None and _is_fio_header(c):
                    idx["fio"] = j
                elif idx["pos"] is None and _is_pos_header(c):
                    idx["pos"] = j
                elif idx["dep"] is None and _is_dep_header(c):
                    idx["dep"] = j
            if idx["fio"] is not None:
                return idx, i
    ncols = max((len(r) for r in rows), default=0)
    stats = {}
    for j in range(ncols):
        s = _stats(rows, j)
        if s and s["digit"] < 0.5:
            stats[j] = s
    if not stats:
        return {"fio": 0, "pos": 1, "dep": 2}, -1
    fio = max(stats.items(),
              key=lambda t: (t[1]["caps"] > 0.6 and t[1]["words"] >= 2, t[1]["card"]))[0]
    rest = [j for j in stats if j != fio]
    dep = None
    kw = [(j, stats[j]["dep_kw"]) for j in rest if stats[j]["dep_kw"] > 0]
    if kw:
        dep = max(kw, key=lambda t: t[1])[0]
    else:
        mid = [j for j in rest if 5 < stats[j]["card"] <= 80]
        if mid:
            dep = min(mid, key=lambda j: stats[j]["card"])
    rest2 = [j for j in rest if j != dep]
    pos = max(rest2, key=lambda j: stats[j]["card"]) if rest2 else None
    return {"fio": fio, "pos": pos, "dep": dep}, -1


def parse_employees(rows):
    idx, header_i = detect_columns(rows)
    data = rows[header_i + 1:] if header_i >= 0 else rows
    emps, seen = [], set()
    for r in data:
        def cell(k):
            j = idx.get(k)
            if j is None or j >= len(r):
                return ""
            return str(r[j]).strip()
        fio = cell("fio")
        if not fio or len(fio) < 3 or re.fullmatch(r"[\d\-\+\. ]+", fio):
            continue
        key = fio.lower()
        if key in seen:
            continue
        seen.add(key)
        dep = cell("dep")
        emps.append({
            "Сотрудник": fio,
            "Должность": cell("pos"),
            "Отдел": dep,
            "Подразделение": dep,
        })
    return emps


def choose_excel():
    path = filedialog.askopenfilename(
        title="Выберите файл HR (Excel или CSV)",
        filetypes=[("Excel/CSV", "*.xlsx *.xls *.csv"), ("Все файлы", "*.*")],
    )
    if not path:
        return None
    try:
        rows = _read_rows(path)
    except Exception as e:
        messagebox.showerror("Ошибка чтения", f"Не удалось прочитать файл:\n{e}")
        return None

    employees = parse_employees(rows)
    if not employees:
        messagebox.showwarning("Внимание", "Сотрудники не найдены. Проверьте, что в файле есть колонка с ФИО.")
        return None

    HR_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(HR_FILE, "w", encoding="utf-8") as f:
        json.dump(employees, f, ensure_ascii=False, indent=2)

    deps = len(set(e["Отдел"] for e in employees if e["Отдел"]))
    messagebox.showinfo("Успех", f"✅ Загружено сотрудников: {len(employees)}\n🏢 Отделов: {deps}")
    return employees


def load_hr_employees():
    if not HR_FILE.exists():
        return []
    try:
        with open(HR_FILE, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
        return data if isinstance(data, list) else []
    except Exception:
        return []


# ================= СВЕРКА С e-SEN =================
def get_hr_name(emp):
    return str(emp.get("Сотрудник", "")).strip()


def get_esen_name(emp):
    return str(emp.get("fio", "")).strip()


def find_best_match(hr_name, esen_list, used_esen):
    for i, esen_emp in enumerate(esen_list):
        if i in used_esen:
            continue
        if names_match(hr_name, get_esen_name(esen_emp)):
            return (i, esen_emp)
    return None


def enrich_esen_with_hr(esen_emp, hr_emp):
    esen_emp["department"] = hr_emp.get("Отдел", "")
    esen_emp["hr_position"] = hr_emp.get("Должность", "")
    esen_emp["hr_status"] = hr_emp.get("Состояние", "")
    esen_emp["hr_hire_date"] = hr_emp.get("Дата приема", "")
    esen_emp["hr_birth_date"] = hr_emp.get("Дата рождения", "")
    esen_emp["matched_with_hr"] = True
    return esen_emp


def compare_employees():
    hr = load_json(HR_FILE)
    esen = load_json(ESEN_FILE)

    matched = []
    only_hr = []
    used_esen = set()

    for hr_emp in hr:
        hr_name = get_hr_name(hr_emp)
        if not hr_name:
            continue

        best = find_best_match(hr_name, esen, used_esen)
        if best:
            index, esen_emp = best
            used_esen.add(index)
            esen[index] = enrich_esen_with_hr(esen_emp, hr_emp)
            matched.append({
                "hr_name": hr_name,
                "esen_name": get_esen_name(esen_emp),
                "department": hr_emp.get("Отдел", ""),
                "hr_position": hr_emp.get("Должность", ""),
                "hr": hr_emp,
                "esen": esen[index],
            })
        else:
            only_hr.append(hr_emp)

    only_esen = []
    for i, esen_emp in enumerate(esen):
        if i not in used_esen:
            esen_emp["matched_with_hr"] = False
            only_esen.append(esen_emp)

    result = {
        "total_hr": len(hr),
        "total_esen": len(esen),
        "matched": matched,
        "possible_matches": [],
        "only_hr": only_hr,
        "only_esen": only_esen,
    }
    result["synced_at"] = datetime.now().isoformat(timespec="seconds")
    append_sync_history(result)

    save_json(ESEN_FILE, esen)
    save_json(RESULT_FILE, result)

    print("✅ Сверка завершена")
    print(f"HR Excel: {len(hr)}")
    print(f"e-SEN: {len(esen)}")
    print(f"Совпали: {len(matched)}")
    print(f"Есть в HR, но нет в e-SEN: {len(only_hr)}")
    print(f"Есть в e-SEN, но нет в HR: {len(only_esen)}")

    export_missing_separate_files_by_departments()
    return result


def run_compare():
    return compare_employees()


def run():
    return compare_employees()
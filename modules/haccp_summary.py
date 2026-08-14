import json
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox

import customtkinter as ctk


PROJECT_ROOT = Path(__file__).resolve().parent.parent
RECORDS_FILE = (
    PROJECT_ROOT
    / "database"
    / "haccp_temperature_records.json"
)

ALL_OBJECTS = "Все объекты"
ALL_DEPARTMENTS = "Все подразделения"
ALL_STATUSES = "Все статусы"


def load_records():
    if not RECORDS_FILE.exists():
        return []

    try:
        with RECORDS_FILE.open(
            "r",
            encoding="utf-8-sig"
        ) as file:
            data = json.load(file)

        return data if isinstance(data, list) else []

    except Exception as error:
        messagebox.showerror(
            "SanEpi AI",
            (
                "Не удалось прочитать "
                f"журнал температур:\n{error}"
            )
        )
        return []


def format_temperature(value):
    try:
        number = float(value)

        if number.is_integer():
            return f"{int(number)} °C"

        return f"{number:g} °C"

    except (TypeError, ValueError):
        return "-"


def export_to_excel(
    records,
    parent
):
    if not records:
        messagebox.showwarning(
            "SanEpi AI",
            "Нет записей для экспорта.",
            parent=parent
        )
        return

    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            Alignment,
            Font,
            PatternFill,
        )
        from openpyxl.utils import (
            get_column_letter,
        )

    except ImportError:
        messagebox.showerror(
            "Требуется openpyxl",
            (
                "Для экспорта в Excel "
                "необходимо установить openpyxl.\n\n"
                "В терминале выполните:\n"
                "pip install openpyxl"
            ),
            parent=parent
        )
        return

    default_name = (
        "HACCP_журнал_температур_"
        + datetime.now().strftime(
            "%Y-%m-%d"
        )
        + ".xlsx"
    )

    output_path = filedialog.asksaveasfilename(
        parent=parent,
        title="Сохранить журнал HACCP",
        defaultextension=".xlsx",
        initialfile=default_name,
        filetypes=[
            (
                "Excel",
                "*.xlsx"
            )
        ]
    )

    if not output_path:
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Журнал температур"

    sheet.merge_cells(
        "A1:L1"
    )

    title_cell = sheet["A1"]
    title_cell.value = (
        "HACCP — журнал контроля температур"
    )
    title_cell.font = Font(
        bold=True,
        size=16,
        color="FFFFFF"
    )
    title_cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1D4ED8"
    )
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    sheet.row_dimensions[1].height = 28

    headers = [
        "Дата",
        "Время",
        "Объект",
        "Подразделение",
        "Оборудование",
        "Температура, °C",
        "Минимум, °C",
        "Максимум, °C",
        "Статус",
        "Ответственный",
        "Корректирующее действие",
        "Дата создания записи",
    ]

    for column, header in enumerate(
        headers,
        start=1
    ):
        cell = sheet.cell(
            row=3,
            column=column,
            value=header
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="2563EB"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    for row_number, record in enumerate(
        records,
        start=4
    ):
        values = [
            record.get("date", ""),
            record.get("time", ""),
            record.get("object_name", ""),
            record.get(
                "department_name",
                ""
            ),
            record.get(
                "equipment_name",
                ""
            ),
            record.get(
                "temperature",
                ""
            ),
            record.get(
                "temperature_min",
                ""
            ),
            record.get(
                "temperature_max",
                ""
            ),
            record.get("status", ""),
            record.get(
                "responsible",
                ""
            ),
            record.get(
                "corrective_action",
                ""
            ),
            record.get(
                "created_at",
                ""
            ),
        ]

        for column, value in enumerate(
            values,
            start=1
        ):
            cell = sheet.cell(
                row=row_number,
                column=column,
                value=value
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

        status = record.get(
            "status",
            ""
        )

        status_cell = sheet.cell(
            row=row_number,
            column=9
        )

        if status == "Отклонение":
            status_cell.fill = PatternFill(
                fill_type="solid",
                fgColor="FECACA"
            )
            status_cell.font = Font(
                bold=True,
                color="991B1B"
            )

        else:
            status_cell.fill = PatternFill(
                fill_type="solid",
                fgColor="DCFCE7"
            )
            status_cell.font = Font(
                bold=True,
                color="166534"
            )

    widths = {
        1: 13,
        2: 10,
        3: 28,
        4: 24,
        5: 25,
        6: 17,
        7: 16,
        8: 16,
        9: 15,
        10: 24,
        11: 42,
        12: 22,
    }

    for column, width in widths.items():
        sheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = (
        f"A3:L{sheet.max_row}"
    )

    try:
        workbook.save(
            output_path
        )

        messagebox.showinfo(
            "SanEpi AI",
            (
                "Журнал успешно сохранён:\n\n"
                f"{output_path}"
            ),
            parent=parent
        )

    except Exception as error:
        messagebox.showerror(
            "Ошибка экспорта",
            str(error),
            parent=parent
        )


def open_haccp_summary(parent):
    records = load_records()

    window = ctk.CTkToplevel(parent)
    window.title("Общий журнал HACCP")
    window.geometry("1250x780")
    window.minsize(1000, 650)
    window.lift()
    window.focus_force()
    window.grab_set()

    ctk.CTkLabel(
        window,
        text="📊 Общий журнал температур",
        font=("Arial", 28, "bold")
    ).pack(
        pady=(20, 12)
    )

    filters_frame = ctk.CTkFrame(
        window,
        corner_radius=12
    )
    filters_frame.pack(
        fill="x",
        padx=20,
        pady=(0, 10)
    )

    filters_frame.grid_columnconfigure(
        0,
        weight=1
    )
    filters_frame.grid_columnconfigure(
        1,
        weight=1
    )
    filters_frame.grid_columnconfigure(
        2,
        weight=1
    )
    filters_frame.grid_columnconfigure(
        3,
        weight=1
    )

    ctk.CTkLabel(
        filters_frame,
        text="Дата",
        font=("Arial", 13, "bold")
    ).grid(
        row=0,
        column=0,
        sticky="w",
        padx=10,
        pady=(10, 4)
    )

    date_entry = ctk.CTkEntry(
        filters_frame,
        height=38,
        placeholder_text=(
            "ГГГГ-ММ-ДД или пусто"
        )
    )
    date_entry.grid(
        row=1,
        column=0,
        sticky="ew",
        padx=10,
        pady=(0, 10)
    )
    date_entry.insert(
        0,
        datetime.now().strftime(
            "%Y-%m-%d"
        )
    )

    object_names = sorted(
        {
            str(
                record.get(
                    "object_name",
                    ""
                )
            ).strip()
            for record in records
            if str(
                record.get(
                    "object_name",
                    ""
                )
            ).strip()
        }
    )

    selected_object = ctk.StringVar(
        value=ALL_OBJECTS
    )

    ctk.CTkLabel(
        filters_frame,
        text="Объект",
        font=("Arial", 13, "bold")
    ).grid(
        row=0,
        column=1,
        sticky="w",
        padx=10,
        pady=(10, 4)
    )

    object_menu = ctk.CTkOptionMenu(
        filters_frame,
        variable=selected_object,
        values=[
            ALL_OBJECTS
        ] + object_names,
        height=38
    )
    object_menu.grid(
        row=1,
        column=1,
        sticky="ew",
        padx=10,
        pady=(0, 10)
    )

    selected_department = ctk.StringVar(
        value=ALL_DEPARTMENTS
    )

    ctk.CTkLabel(
        filters_frame,
        text="Подразделение",
        font=("Arial", 13, "bold")
    ).grid(
        row=0,
        column=2,
        sticky="w",
        padx=10,
        pady=(10, 4)
    )

    department_menu = ctk.CTkOptionMenu(
        filters_frame,
        variable=selected_department,
        values=[ALL_DEPARTMENTS],
        height=38
    )
    department_menu.grid(
        row=1,
        column=2,
        sticky="ew",
        padx=10,
        pady=(0, 10)
    )

    selected_status = ctk.StringVar(
        value=ALL_STATUSES
    )

    ctk.CTkLabel(
        filters_frame,
        text="Статус",
        font=("Arial", 13, "bold")
    ).grid(
        row=0,
        column=3,
        sticky="w",
        padx=10,
        pady=(10, 4)
    )

    status_menu = ctk.CTkOptionMenu(
        filters_frame,
        variable=selected_status,
        values=[
            ALL_STATUSES,
            "Норма",
            "Отклонение",
        ],
        height=38
    )
    status_menu.grid(
        row=1,
        column=3,
        sticky="ew",
        padx=10,
        pady=(0, 10)
    )

    action_frame = ctk.CTkFrame(
        window,
        fg_color="transparent"
    )
    action_frame.pack(
        fill="x",
        padx=20,
        pady=(0, 10)
    )

    summary_label = ctk.CTkLabel(
        action_frame,
        text="",
        font=("Arial", 15, "bold")
    )
    summary_label.pack(
        side="left"
    )

    export_button = ctk.CTkButton(
        action_frame,
        text="📥 Экспорт в Excel",
        width=190,
        height=38,
        fg_color="#059669",
        hover_color="#047857"
    )
    export_button.pack(
        side="right"
    )

    refresh_button = ctk.CTkButton(
        action_frame,
        text="🔄 Применить фильтры",
        width=190,
        height=38
    )
    refresh_button.pack(
        side="right",
        padx=(0, 8)
    )

    records_frame = ctk.CTkScrollableFrame(
        window,
        corner_radius=12
    )
    records_frame.pack(
        fill="both",
        expand=True,
        padx=20,
        pady=(0, 20)
    )

    filtered_records = {
        "value": []
    }

    def update_departments(
        _selected_value=None
    ):
        object_filter = selected_object.get()

        departments = sorted(
            {
                str(
                    record.get(
                        "department_name",
                        ""
                    )
                ).strip()
                for record in records
                if str(
                    record.get(
                        "department_name",
                        ""
                    )
                ).strip()
                and (
                    object_filter == ALL_OBJECTS
                    or record.get(
                        "object_name"
                    ) == object_filter
                )
            }
        )

        department_menu.configure(
            values=[
                ALL_DEPARTMENTS
            ] + departments
        )

        selected_department.set(
            ALL_DEPARTMENTS
        )

    def get_filtered_records():
        date_filter = (
            date_entry.get().strip()
        )

        object_filter = (
            selected_object.get()
        )

        department_filter = (
            selected_department.get()
        )

        status_filter = (
            selected_status.get()
        )

        result = []

        for record in records:
            if (
                date_filter
                and record.get("date")
                != date_filter
            ):
                continue

            if (
                object_filter
                != ALL_OBJECTS
                and record.get(
                    "object_name"
                ) != object_filter
            ):
                continue

            if (
                department_filter
                != ALL_DEPARTMENTS
                and record.get(
                    "department_name"
                ) != department_filter
            ):
                continue

            if (
                status_filter
                != ALL_STATUSES
                and record.get(
                    "status"
                ) != status_filter
            ):
                continue

            result.append(record)

        result.sort(
            key=lambda item: (
                item.get("date", ""),
                item.get("time", ""),
            ),
            reverse=True
        )

        return result

    def render_records():
        for widget in (
            records_frame.winfo_children()
        ):
            widget.destroy()

        result = get_filtered_records()

        filtered_records["value"] = result

        normal_count = sum(
            1
            for record in result
            if record.get("status")
            == "Норма"
        )

        deviation_count = sum(
            1
            for record in result
            if record.get("status")
            == "Отклонение"
        )

        summary_label.configure(
            text=(
                f"Всего: {len(result)}   "
                f"✅ Норма: {normal_count}   "
                f"❌ Отклонения: "
                f"{deviation_count}"
            )
        )

        if not result:
            ctk.CTkLabel(
                records_frame,
                text=(
                    "По выбранным фильтрам "
                    "записи не найдены."
                ),
                font=("Arial", 16),
                text_color="#9ca3af"
            ).pack(
                pady=70
            )
            return

        for record in result:
            status = record.get(
                "status",
                ""
            )

            is_deviation = (
                status == "Отклонение"
            )

            card = ctk.CTkFrame(
                records_frame,
                corner_radius=10,
                border_width=2,
                border_color=(
                    "#ef4444"
                    if is_deviation
                    else "#22c55e"
                )
            )
            card.pack(
                fill="x",
                padx=7,
                pady=6
            )

            title = (
                f"{record.get('date', '')} "
                f"{record.get('time', '')}  |  "
                f"{record.get('object_name', '')}  |  "
                f"{record.get('department_name', '')}"
            )

            ctk.CTkLabel(
                card,
                text=title,
                font=("Arial", 16, "bold"),
                anchor="w"
            ).pack(
                fill="x",
                padx=14,
                pady=(12, 4)
            )

            details = (
                f"Оборудование: "
                f"{record.get('equipment_name', '')}\n"
                f"Температура: "
                f"{format_temperature(record.get('temperature'))}\n"
                f"Статус: {status}\n"
                f"Ответственный: "
                f"{record.get('responsible', '')}"
            )

            corrective = str(
                record.get(
                    "corrective_action",
                    ""
                )
            ).strip()

            if corrective:
                details += (
                    "\nКорректирующее действие: "
                    f"{corrective}"
                )

            ctk.CTkLabel(
                card,
                text=details,
                font=("Arial", 14),
                anchor="w",
                justify="left",
                wraplength=950
            ).pack(
                fill="x",
                padx=14,
                pady=(4, 12)
            )

    def export_current_records():
        export_to_excel(
            records=filtered_records["value"],
            parent=window
        )

    object_menu.configure(
        command=update_departments
    )

    refresh_button.configure(
        command=render_records
    )

    export_button.configure(
        command=export_current_records
    )

    update_departments()
    render_records()
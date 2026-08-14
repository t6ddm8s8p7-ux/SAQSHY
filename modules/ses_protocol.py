"""
Протокол (программа) производственного контроля по приказу № 62 от 07.04.2023 (приложение 1).
Просмотр + выгрузка в PDF и Excel для каждого объекта СЭС.
"""
import os
import textwrap
from datetime import date
from pathlib import Path
from tkinter import messagebox
import customtkinter as ctk

from modules.ses_lab_data import find_lab_categories

PROJECT_ROOT = Path(__file__).resolve().parent.parent
EXPORTS_DIR = PROJECT_ROOT / "exports" / "ses"

ORDER_PK = "Приказ № 62 от 07.04.2023, приложение 1"

# (наименование исследования, определяемые показатели, место отбора проб, периодичность)
PROTOCOL_ROWS = {
    "catering": [
        ("Блюда из мяса, птицы, рыбы (выборочно)", "микробиологические по ТР ТС 021/2011", "кухня, рестораны, бары", "1 раз в квартал – 1 раз в год (по типу объекта)"),
        ("Гарниры, салаты с заправками, напитки собств. производства", "микробиологические по ТР ТС 021/2011", "кухня, бары", "1 раз в квартал – 1 раз в год"),
        ("Смывы с инвентаря, рук, оборудования, сан одежды", "БГКП", "кухня, бары", "5–20 смывов, 1–4 раза в год"),
        ("Готовая продукция после термообработки", "качество тепловой обработки", "кухня", "1–2 раза в год"),
        ("Вода с водопроводной сети", "бактериологические, краткий хим анализ", "водоразборные краны", "1 раз в год"),
        ("Дезинфицирующие растворы", "АДВ (% активности)", "кухня, бары", "1 раз в квартал"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в полгода"),
    ],
    "food_production": [
        ("Сырьё", "органолептика, физико-химия, паразитология", "склады, производство", "ежеквартально, 1 проба"),
        ("Яйцо", "овоскопирование", "каждая партия", "10 яиц от партии"),
        ("Готовая продукция", "микробиология, токсичные элементы, микотоксины", "каждая партия", "1 раз в год"),
        ("Смывы с чистого инвентаря, тары, рук", "БГКП, патогенная микрофлора", "цеха", "ежеквартально, 10 смывов"),
        ("Технологический процесс", "температурно-временные режимы", "каждый цикл", "каждый технологический цикл"),
        ("Вода питьевая", "органолептика, физико-химия, микробиология", "краны", "2 раза в год (местные источники)"),
        ("Производственные и складские помещения", "температура, относительная влажность", "все помещения", "1 раз в неделю"),
        ("Дезинфицирующие растворы", "АДВ", "производство", "1 раз в квартал, 2 пробы"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в полгода"),
    ],
    "pool_spa": [
        ("Вода чаши бассейна", "микробиология (ОМЧ, ОКБ, ТКБ, стафилококк, синегнойная палочка, легионелла)", "2 точки: мелкая и глубокая часть", "1 раз в месяц, 2 пробы"),
        ("Вода чаши бассейна", "сан-хим (остаточный хлор, pH, мутность, температура)", "чаша бассейна", "в рабочие часы / каждые 4 ч / 1 раз в квартал"),
        ("Вода чаши бассейна", "паразитологические (цисты лямблий, яйца гельминтов)", "чаша бассейна", "2 раза в год, 3–5 проб"),
        ("Смывы со стенок, скамеек, рук, спецодежды", "БГКП", "душевые, раздевалки", "2 раза в год, 10–15 проб"),
        ("Воздух и микроклимат зала бассейна", "температура, влажность, скорость движения воздуха", "зал бассейна", "1 раз в рабочие часы"),
        ("Текущая уборка и дезинфекция помещений", "смывы, дезсредства и растворы", "помещения", "1 раз в квартал, 10 смывов, 2 пробы"),
        ("Вентиляция", "кратность воздухообмена", "зал бассейна", "1 раз в год / 1 раз в 3 года"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в полгода"),
    ],
    "cosmetology": [
        ("Обработка инструментария (маникюр/педикюр/косметология)", "качество стерилизации, предстерилизационная очистка", "кабинеты", "1 раз в квартал"),
        ("Вода", "краткий хим анализ + микробиология", "краны", "при вводе, после реконструкции"),
        ("Микроклимат производственных помещений", "температура, влажность", "кабинеты", "1 раз в год"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в полгода"),
    ],
    "hotel": [
        ("Постельные принадлежности (матрацы, подушки, одеяла)", "эффективность камерной обработки", "после обработки", "1 раз в полгода, не менее 1%"),
        ("Микроклимат номеров и рабочих мест", "температура, влажность", "номера, рабочие места", "1 раз в год"),
        ("Вода", "бактериологические, сан-хим", "краны", "при вводе, после ремонта, после аварий"),
        ("Эффективность дезинсекции/дератизации", "насекомые, грызуны", "объект", "1 раз в полугодие"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в полгода"),
    ],
    "water_supply": [
        ("Питьевая вода", "микробиология, сан-хим, паразитология, радиология", "водозабор, перед подачей в сеть, сеть", "по программе (приказ № 26)"),
        ("Дезинфицирующие средства", "АДВ", "рабочие растворы", "1 раз в квартал"),
        ("Наблюдательные скважины", "органолептика, хим, микробиология, радиология", "скважины", "2 раза в год"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "2 раза в год"),
    ],
    "sewage": [
        ("Сточные воды при выпуске", "органолептика, сан-хим, микробиология", "выпуск", "1 раз в квартал"),
        ("Почва при сбросе на почву", "хим, микробиология, паразитология", "почва", "1 раз в квартал"),
        ("Атмосферный воздух на границе СЗЗ", "хим показатели", "граница СЗЗ", "по проекту"),
        ("Вентиляция", "кратность воздухообмена", "сооружения", "1 раз в год / 1 раз в 3 года"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в полгода"),
    ],
    "healthcare": [
        ("Бактериальная обсеменённость воздуха", "золотистый стафилококк, грибы", "процедурные, стерилизационные", "1 раз в квартал"),
        ("Смывы с внешней среды", "БГКП, патогенная флора", "кабинеты, инвентарь, руки", "1 раз в квартал"),
        ("Исследование на стерильность", "стерильность", "инструментарий", "1 раз в месяц"),
        ("Предстерилизационная очистка", "скрытая кровь, щёлочь", "ЦСО", "ежедневно"),
        ("Дезинфицирующие средства", "АДВ", "кабинеты", "1 раз в квартал"),
        ("Микроклимат", "температура, влажность", "палаты, кабинеты", "2 раза в год"),
        ("Вентиляция", "кратность воздухообмена", "кабинеты", "1 раз в год / 1 раз в 3 года"),
        ("Освещённость", "уровни освещённости", "кабинеты", "1 раз в год"),
        ("Вода", "бактериологические, сан-хим", "краны", "при вводе, после ремонта, после аварий"),
        ("Индивидуальный дозиметрический контроль персонала группы А", "уровень радиации", "кабинеты с ИИИ", "1 раз в квартал"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в полгода"),
    ],
    "lab": [
        ("Воздух (боксы, стерилизационные)", "КОЕ, золотистый стафилококк, грибы", "боксы, стерилизационные", "1 раз в квартал"),
        ("Смывы", "БГКП, патогенная флора", "оборудование, руки", "1 раз в квартал"),
        ("Стерильность", "стерильность", "инструментарий, посуда", "1 раз в месяц"),
        ("Дезинфицирующие средства", "АДВ", "лаборатории", "1 раз в квартал"),
        ("Микроклимат", "температура, влажность", "рабочие места", "2 раза в год"),
        ("Вентиляция", "кратность воздухообмена", "лаборатории", "1 раз в год / 1 раз в 3 года"),
        ("Освещённость", "уровни освещённости", "рабочие места", "1 раз в год"),
        ("Вода", "бактериологические, сан-хим", "краны", "1 раз в год"),
    ],
    "radiation": [
        ("Индивидуальный дозиметрический контроль персонала группы А", "уровень радиации", "рабочие места", "1 раз в квартал"),
        ("МЭД гамма-излучения", "уровень радиации", "рабочие места, смежные помещения", "не реже 1 раза в год"),
        ("Эффективность средств защиты", "толщина свинцового слоя", "средства защиты", "не реже 1 раза в 2 года"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в год"),
    ],
    "disinf": [
        ("Рабочие растворы дезсредств", "концентрация АДВ", "рабочие помещения", "1 раз в квартал"),
        ("Микроклимат склада", "температура, влажность", "склад", "2 раза в год"),
        ("Воздух рабочей зоны", "содержание АДВ", "производство, склад", "2 раза в год"),
        ("Освещённость", "уровни освещённости", "склад", "1 раз в год"),
        ("Вентиляция", "кратность воздухообмена", "помещения", "1 раз в год / 1 раз в 3 года"),
        ("Вода", "бактериологические, сан-хим", "краны", "1 раз в год"),
    ],
    "transport": [
        ("Вода с системы водоснабжения", "бактериологические, сан-хим", "ёмкости, краны", "1 раз в полгода"),
        ("Микроклимат", "температура, влажность, скорость воздуха", "салоны, кабины", "1 раз в полгода"),
        ("Постельные принадлежности", "смывы на яйца гельминтов", "салоны", "1 раз в полгода, 10 смывов"),
        ("Дезинфицирующие растворы", "АДВ", "транспорт", "1 раз в полгода"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в год"),
    ],
    "market": [
        ("Смывы с оборудования, инвентаря, рук", "БГКП", "торговые места", "1 раз в полгода, 10 смывов"),
        ("Молочная, колбасная продукция, птица, рыба, кондитерские изделия", "микробиологические по ТР ТС 021/2011", "прилавки, склады", "1 раз в полгода, 1–2 пробы"),
        ("Яйцо куриное", "микробиологические", "торговля", "1 раз в полгода"),
        ("Эффективность дезинсекции/дератизации", "насекомые, грызуны", "объект", "1 раз в полугодие"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в полгода"),
    ],
    "storage": [
        ("Смывы с оборудования, инвентаря, рук", "БГКП", "склады", "1 раз в полгода, 10 смывов"),
        ("Молочная, колбасная продукция, птица, рыба, кондитерские изделия", "микробиологические по ТР ТС 021/2011", "склады", "1 раз в полгода, 2 пробы"),
        ("Микроклимат", "температура, влажность", "склады", "согласно режиму хранения"),
        ("Дезинфицирующие растворы", "АДВ", "склады", "1 раз в полгода"),
        ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в полгода"),
    ],
}

DEFAULT_ROWS = [
    ("Смывы с инвентаря, рук, оборудования", "БГКП", "производственные помещения", "1 раз в полгода, 10 смывов"),
    ("Питьевая вода", "бактериологические, сан-хим", "краны", "1 раз в год"),
    ("Микроклимат", "температура, влажность", "рабочие места", "2 раза в год"),
    ("Вентиляция", "кратность воздухообмена", "помещения", "1 раз в год / 1 раз в 3 года"),
    ("Освещённость", "уровни освещённости", "рабочие места", "1 раз в год"),
    ("Дезинфицирующие растворы", "АДВ", "рабочие растворы", "1 раз в квартал"),
    ("Эффективность дезинсекции/дератизации", "насекомые, грызуны", "объект", "1 раз в полугодие"),
    ("Обязательные медосмотры", "полнота, своевременность", "персонал", "1 раз в полгода"),
]


def get_protocol_rows(object_name):
    rows = []
    for cat in find_lab_categories(object_name):
        for study, indicators, point, freq in PROTOCOL_ROWS.get(cat["id"], []):
            rows.append((cat["title"], study, indicators, point, freq))
    if not rows:
        for study, indicators, point, freq in DEFAULT_ROWS:
            rows.append(("Базовый набор", study, indicators, point, freq))
    return rows


def _safe_name(value):
    name = str(value or "объект").strip()
    for ch in '\\/:*?"<>|':
        name = name.replace(ch, "-")
    return name[:60]


def _get_fonts():
    candidates = [
        (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
        (r"C:\Windows\Fonts\calibri.ttf", r"C:\Windows\Fonts\calibrib.ttf"),
    ]
    for reg, bold in candidates:
        if Path(reg).exists():
            return reg, (bold if Path(bold).exists() else reg)
    return None, None


def export_protocol_pdf(obj):
    """Выгружает протокол производственного контроля в PDF."""
    import fitz
    name = str(obj.get("name", "Объект"))
    category = str(obj.get("category", "-"))
    rows = get_protocol_rows(name)

    font_reg, font_bold = _get_fonts()
    if not font_reg:
        messagebox.showerror("SanEpi AI", "Не найден шрифт Arial/Calibri в Windows.")
        return None

    doc = fitz.open()
    W, H, M = 595, 842, 40
    page = doc.new_page(width=W, height=H)
    state = {"y": M}

    def new_page():
        nonlocal page
        page = doc.new_page(width=W, height=H)
        state["y"] = M

    def write(text, size=10, bold=False, gap=3):
        width_chars = max(30, int((W - 2 * M) / (size * 0.55)))
        lines = textwrap.wrap(str(text), width=width_chars) or [""]
        for ln in lines:
            if state["y"] > H - M:
                new_page()
            page.insert_text(
                fitz.Point(M, state["y"]),
                ln,
                fontname="fb" if bold else "fr",
                fontfile=font_bold if bold else font_reg,
                fontsize=size,
            )
            state["y"] += size * 1.35
        state["y"] += gap

    write("ПРОТОКОЛ (ПРОГРАММА) ПРОИЗВОДСТВЕННОГО КОНТРОЛЯ", size=13, bold=True, gap=8)
    write(f"Объект: {name}", size=11, bold=True)
    write(f"Эпидемическая значимость: {category}", size=10)
    write(f"Основание: {ORDER_PK}; Перечень — Приказ ҚР ДСМ-220/2020", size=9)
    write(f"Дата формирования: {date.today().isoformat()}", size=9, gap=10)

    number = 0
    current_cat = None
    for cat, study, indicators, point, freq in rows:
        if cat != current_cat:
            current_cat = cat
            write(f"■ {cat}", size=11, bold=True, gap=4)
        number += 1
        write(f"{number}. {study}", size=10, bold=True)
        write(f"Определяемые показатели: {indicators}", size=9, indent=0)
        write(f"Место отбора проб / замеров: {point}", size=9)
        write(f"Периодичность (не менее): {freq}", size=9, gap=7)

    write("", size=6)
    write("Ответственный за производственный контроль: ______________________", size=10)
    write("Руководитель объекта: ______________________", size=10)

    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    path = EXPORTS_DIR / f"PK_{_safe_name(name)}_{date.today()}.pdf"
    try:
        doc.save(str(path))
        doc.close()
    except Exception as e:
        messagebox.showerror("SanEpi AI", f"Не удалось сохранить PDF:\n{e}")
        return None
    messagebox.showinfo("SanEpi AI", f"✅ Протокол ПК (PDF) сохранён:\n\n{path}")
    try:
        os.startfile(str(path))
    except Exception:
        pass
    return path


def export_protocol_excel(obj):
    name = str(obj.get("name", "Объект"))
    category = str(obj.get("category", "-"))
    rows = get_protocol_rows(name)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        messagebox.showerror("SanEpi AI", "Требуется openpyxl:\npip install openpyxl")
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "Протокол ПК"
    thin = Side(style="thin", color="9CA3AF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1D4ED8")
    wrap = Alignment(vertical="top", wrap_text=True)
    center = Alignment(horizontal="center", vertical="top", wrap_text=True)

    ws.merge_cells("A1:E1")
    t = ws["A1"]
    t.value = "ПРОГРАММА (ПРОТОКОЛ) ПРОИЗВОДСТВЕННОГО КОНТРОЛЯ"
    t.font = Font(bold=True, size=13, color="FFFFFF")
    t.fill = head_fill
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    info = [
        ("Объект:", name),
        ("Эпидемическая значимость:", category),
        ("Основание:", f"{ORDER_PK}; Перечень — Приказ ҚР ДСМ-220/2020"),
        ("Дата формирования:", date.today().isoformat()),
    ]
    r = 2
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1

    headers = ["№", "Наименование исследований", "Определяемые показатели", "Место отбора проб / замеров", "Периодичность (не менее)"]
    header_row = r
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.border = border
        cell.alignment = center
    r += 1
    for i, (cat, study, indicators, point, freq) in enumerate(rows, 1):
        values = [i, f"{study} ({cat})", indicators, point, freq]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = center if c == 1 else wrap
        r += 1
    widths = {1: 5, 2: 45, 3: 40, 4: 34, 5: 26}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    path = EXPORTS_DIR / f"PK_{_safe_name(name)}_{date.today()}.xlsx"
    try:
        wb.save(path)
    except Exception as e:
        messagebox.showerror("SanEpi AI", f"Не удалось сохранить:\n{e}")
        return None
    messagebox.showinfo("SanEpi AI", f"✅ Протокол ПК (Excel) сохранён:\n\n{path}")
    try:
        os.startfile(str(EXPORTS_DIR))
    except Exception:
        pass
    return path


def open_protocol_window(parent, obj):
    name = str(obj.get("name", "Объект"))
    category = str(obj.get("category", "-"))
    rows = get_protocol_rows(name)

    window = ctk.CTkToplevel(parent)
    window.title("📄 Протокол производственного контроля")
    window.geometry("1200x760")
    window.minsize(980, 620)
    window.lift()
    window.focus_force()

    ctk.CTkLabel(window, text="📄 Протокол производственного контроля", font=("Arial", 24, "bold")).pack(pady=(16, 4))
    ctk.CTkLabel(
        window,
        text=f"{name}  |  {category}\nОснование: {ORDER_PK}",
        font=("Arial", 13),
        text_color="#9ca3af",
        justify="center",
    ).pack(padx=20, pady=(0, 10))

    scroll = ctk.CTkScrollableFrame(window, corner_radius=12)
    scroll.pack(fill="both", expand=True, padx=16, pady=(0, 10))
    scroll.grid_columnconfigure(1, weight=2)
    scroll.grid_columnconfigure(2, weight=2)
    scroll.grid_columnconfigure(3, weight=1)
    scroll.grid_columnconfigure(4, weight=1)

    headers = ["№", "Наименование исследований", "Определяемые показатели", "Место отбора проб", "Периодичность"]
    for c, h in enumerate(headers):
        ctk.CTkLabel(scroll, text=h, font=("Arial", 13, "bold")).grid(row=0, column=c, padx=8, pady=8, sticky="w")
    for i, (cat, study, indicators, point, freq) in enumerate(rows, 1):
        values = [i, study, indicators, point, freq]
        for c, v in enumerate(values):
            ctk.CTkLabel(
                scroll, text=str(v), font=("Arial", 12), anchor="w", justify="left",
                wraplength=330 if c in (1, 2) else 200,
            ).grid(row=i, column=c, padx=8, pady=4, sticky="new")

    buttons = ctk.CTkFrame(window, fg_color="transparent")
    buttons.pack(fill="x", padx=16, pady=(0, 14))
    ctk.CTkButton(
        buttons, text="📥 Скачать PDF", height=42,
        fg_color="#059669", hover_color="#047857",
        command=lambda: export_protocol_pdf(obj),
    ).pack(side="left", fill="x", expand=True, padx=(0, 6))
    ctk.CTkButton(
        buttons, text="📊 Excel", height=42,
        fg_color="#1d4ed8", hover_color="#1e40af",
        command=lambda: export_protocol_excel(obj),
    ).pack(side="left", fill="x", expand=True, padx=(0, 6))
    ctk.CTkButton(
        buttons, text="Закрыть", height=42,
        fg_color="#6b7280", hover_color="#4b5563",
        command=window.destroy,
    ).pack(side="left", fill="x", expand=True)
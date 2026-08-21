# -*- coding: utf-8 -*-
"""Загрузка HR Excel: понимает строки-отделы и делит сотрудников по отделам."""
import json
import re
from pathlib import Path

from tkinter import filedialog, messagebox

PROJECT_ROOT = Path(__file__).resolve().parent.parent
HR_FILE = PROJECT_ROOT / "database" / "hr_employees.json"

DEP_KEYS = ("отдел", "подраздел", "департамент", "служба", "сектор", "цех", "управлен", "бөлім", "department")


def norm(s):
    return str(s if s is not None else "").strip().lower().replace("ё", "е")


def _is_num(v):
    return bool(re.fullmatch(r"[\d\-\+\. ]+", v))


def _is_dash(v):
    return bool(re.fullmatch(r"[\-\—\_\s]+", v))


def _looks_fio(v):
    words = v.split()
    return len(words) >= 2 and not _is_num(v) and sum(1 for w in words if w[:1].isupper()) >= 2


def _read_rows(path):
    path = str(path)
    if path.lower().endswith(".csv"):
        import csv
        with open(path, "r", encoding="utf-8-sig", errors="ignore") as f:
            return [row for row in csv.reader(f)]
    try:
        import pandas as pd
        return pd.read_excel(path, header=None, dtype=str).fillna("").values.tolist()
    except ImportError:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows


def _dept_row_value(row):
    """Строка-отдел: вся строка заполнена одним названием (или одна ячейка со словом «отдел»)."""
    vals = [str(c).strip() for c in row if str(c).strip()]
    if not vals or all(_is_dash(v) for v in vals):
        return None
    if len(set(v.lower() for v in vals)) == 1 and len(vals) >= 2:
        return vals[0]
    if len(vals) == 1 and len(vals[0].split()) >= 2 and any(k in vals[0].lower() for k in DEP_KEYS):
        return vals[0]
    return None


def _is_fio_header(c):
    return ("фио" in c or "сотрудник" in c or "фамилия" in c or "ф.и.о" in c
            or "тегі" in c or "аты" in c or "қызметкер" in c or "full name" in c)


def _is_pos_header(c):
    return ("должност" in c or "лауазым" in c or "позици" in c or "position" in c)


def _is_dep_header(c):
    return any(k in c for k in DEP_KEYS)


def _find_header(rows):
    for i, row in enumerate(rows[:10]):
        cells = [norm(c) for c in row]
        if any(_is_fio_header(c) for c in cells):
            idx = {"fio": None, "pos": None, "dep": None}
            for j, c in enumerate(cells):
                if not c:
                    continue
                if idx["fio"] is None and _is_fio_header(c):
                    idx["fio"] = j
                elif idx["pos"] is None and _is_pos_header(c):
                    idx["pos"] = j
                elif idx["dep"] is None and _is_dep_header(c):
                    idx["dep"] = j
            if idx["fio"] is not None:
                return idx, i
    return None, -1


def parse_employees(rows):
    emps, seen = [], set()

    # ----- формат с шапкой (ФИО / Должность / Отдел) -----
    idx, header_i = _find_header(rows)
    if header_i >= 0:
        for r in rows[header_i + 1:]:
            def cell(k):
                j = idx.get(k)
                if j is None or j >= len(r):
                    return ""
                return str(r[j]).strip()
            fio = cell("fio")
            if not fio or not _looks_fio(fio):
                continue
            key = fio.lower()
            if key in seen:
                continue
            seen.add(key)
            dep = cell("dep")
            emps.append({"Сотрудник": fio, "Должность": cell("pos"),
                         "Отдел": dep, "Подразделение": dep})
        return emps

    # ----- формат со строками-отделами (ваш Excel) -----
    current = ""
    for row in rows:
        dep = _dept_row_value(row)
        if dep:
            current = dep
            continue
        texts = [(j, str(c).strip()) for j, c in enumerate(row) if str(c).strip()]
        if not texts:
            continue
        fio_j = None
        for j, v in texts:
            if _looks_fio(v):
                fio_j = j
                break
        if fio_j is None:
            continue
        fio = dict(texts)[fio_j]
        key = fio.lower()
        if key in seen:
            continue
        seen.add(key)

        pos, state, iin = "", "", ""
        for j, v in texts:
            if j == fio_j:
                continue
            if _is_num(v):
                if len(v) >= 12 and not iin:
                    iin = v
                continue
            if not pos:
                pos = v
            elif not state:
                state = v
        emps.append({"Сотрудник": fio, "Должность": pos,
                     "Отдел": current, "Подразделение": current,
                     "Состояние": state, "ИИН": iin})
    return emps


def choose_excel():
    path = filedialog.askopenfilename(
        title="Выберите файл HR (Excel или CSV)",
        filetypes=[("Excel/CSV", "*.xlsx *.xls *.csv"), ("Все файлы", "*.*")],
    )
    if not path:
        return None
    try:
        rows = _read_rows(path)
    except Exception as e:
        messagebox.showerror("Ошибка чтения", f"Не удалось прочитать файл:\n{e}")
        return None

    employees = parse_employees(rows)
    if not employees:
        messagebox.showwarning("Внимание", "Сотрудники не найдены. Проверьте, что в файле есть ФИО.")
        return None

    HR_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(HR_FILE, "w", encoding="utf-8") as f:
        json.dump(employees, f, ensure_ascii=False, indent=2)

    deps = len(set(e["Отдел"] for e in employees if e["Отдел"]))
    messagebox.showinfo("Успех", f"✅ Загружено сотрудников: {len(employees)}\n🏢 Отделов: {deps}")
    return employees


def load_hr_employees():
    if not HR_FILE.exists():
        return []
    try:
        with open(HR_FILE, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
        return data if isinstance(data, list) else []
    except Exception:
        return []
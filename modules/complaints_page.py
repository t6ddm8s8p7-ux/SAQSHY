# -*- coding: utf-8 -*-
"""Журнал жалоб и обращений (гости, онлайн, СЭС, внутренние)."""
import sqlite3
from datetime import date, timedelta
from pathlib import Path

import customtkinter as ctk
from tkinter import filedialog, messagebox

from modules.translations import get_language
from modules.complaint_translations import translate_list, get_complaint_translation

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DB_PATH = PROJECT_ROOT / "database" / "sanepi.db"

SOURCES = ["Гость (устно)", "Гость (книга жалоб)", "Комментарии гостей (карточки отеля)",
           "Онлайн-отзыв (2GIS/Google)", "Booking/Expedia",
           "СЭС / ДСЭК", "Прокуратура / Акимат", "Внутренний аудит", "Соцсети (Instagram)"]
CATEGORIES = ["Питание /Restaurant", "Чистота / Номера", "Обслуживание", "Бассейн / Аквапарк",
              "СПА / Массаж", "Шум / Кондиционер", "Безопасность", "Детский клуб", "Прачечная",
              "Трансфер / Парковка", "Wi-Fi / Техника", "Другое"]
STATUSES = ["🆕 Новая", "⚙️ В работе", "✅ Решена", "❌ Отклонена", " Повторная"]
PRIORITIES = ["🔴 Высокая", "🟡 Средняя", "🟢 Низкая"]

COMPLAINT_TYPES = [
    "Пищевое отравление / подозрение на отравление",
    "Некачественная еда / холодное блюдо",
    "Медленное обслуживание в ресторане",
    "Грязный номер / плохая уборка",
    "Неприятный запах в номере / помещении",
    "Насекомые / грызуны в помещении",
    "Бассейн: температура / чистота воды",
    "Шум от соседей / мероприятий",
    "Не работает кондиционер",
    "Проблема с Wi-Fi / техникой",
    "Грубость персонала",
    "Долгое заселение / очередь на ресепшн",
    "СПА: замечания по записи / обслуживанию",
    "Детский клуб: замечания",
    "Трансфер / парковка: замечания",
    "Безопасность: замечание",
]

T = {
    "title": {"ru": "📢 Журнал жалоб", "kk": "📢 Шағымдар журналы", "en": "📢 Complaints log", "tr": "📢 Şikayet günlüğü"},
    "add": {"ru": "➕ Новая", "kk": "➕ Жаңа", "en": "➕ New", "tr": "➕ Yeni"},
    "export": {"ru": "📥 Excel", "kk": "📥 Excel", "en": " Excel", "tr": "📥 Excel"},
    "no_rec": {"ru": "Жалоб за период нет.", "kk": "Кезеңде шағымдар жоқ.", "en": "No complaints.", "tr": "Bu dönemde şikayet yok."},
    "warning": {"ru": "⚠️ Внимание", "kk": "⚠️ Назар аударыңыз", "en": "⚠️ Warning", "tr": "⚠️ Dikkat"},
    "del": {"ru": "Удалить жалобу?", "kk": "Шағымды жою?", "en": "Delete complaint?", "tr": "Şikayet silinsin mi?"},
    "need_text": {"ru": "Опишите суть жалобы!", "kk": "Шағымның мәнін жазыңыз!", "en": "Describe the complaint!", "tr": "Şikayeti açıklayın!"},
    "from": {"ru": "с", "kk": "бастап", "en": "from", "tr": "başlangıç"},
    "to": {"ru": "по", "kk": "дейін", "en": "to", "tr": "bitiş"},
    "save": {"ru": "💾 Сохранить", "kk": "💾 Сақтау", "en": " Save", "tr": "💾 Kaydet"},
    "date": {"ru": "Дата:", "kk": "Күні:", "en": "Date:", "tr": "Tarih:"},
    "source": {"ru": "Источник:", "kk": "Көзі:", "en": "Source:", "tr": "Kaynak:"},
    "source_note": {"ru": "💬 Комментарий (ссылка, № комнаты):", "kk": "💬 Түсініктеме (сілтеме, № бөлме):", "en": "💬 Comment (link, room №):", "tr": "💬 Yorum (link, oda no):"},
    "category": {"ru": "Категория:", "kk": "Санат:", "en": "Category:", "tr": "Kategori:"},
    "priority": {"ru": "Приоритет:", "kk": "Басымдық:", "en": "Priority:", "tr": "Öncelik:"},
    "status": {"ru": "Статус:", "kk": "Мәртебе:", "en": "Status:", "tr": "Durum:"},
    "guest": {"ru": "Гость / № комнаты:", "kk": "Қонақ / № бөлме:", "en": "Guest / Room:", "tr": "Misafir / Oda:"},
    "department": {"ru": "Отдел:", "kk": "Бөлім:", "en": "Department:", "tr": "Departman:"},
    "assignee": {"ru": "Исполнитель:", "kk": "Орындаушы:", "en": "Assignee:", "tr": "Sorumlu:"},
    "deadline": {"ru": "Срок:", "kk": "Мерзімі:", "en": "Deadline:", "tr": "Son tarih:"},
    "description": {"ru": "Суть жалобы:", "kk": "Шағымның мәні:", "en": "Complaint details:", "tr": "Şikayet detayı:"},
    "measures": {"ru": "Принятые меры:", "kk": "Қабылданған шаралар:", "en": "Measures taken:", "tr": "Alınan önlemler:"},
    "result": {"ru": "Итог / ответ:", "kk": "Нәтиже / жауап:", "en": "Result / reply:", "tr": "Sonuç / yanıt:"},
    "all": {"ru": "Все", "kk": "Барлығы", "en": "All", "tr": "Tümü"},
    "stat_total": {"ru": "Всего", "kk": "Барлығы", "en": "Total", "tr": "Toplam"},
    "stat_new": {"ru": "🆕 Новых", "kk": "🆕 Жаңа", "en": "🆕 New", "tr": " Yeni"},
    "stat_work": {"ru": "⚙️ В работе", "kk": "⚙️ Жұмыста", "en": "⚙️ In progress", "tr": "⚙️ Devam eden"},
    "stat_done": {"ru": "✅ Решено", "kk": "✅ Шешілді", "en": "✅ Resolved", "tr": "✅ Çözüldü"},
    "stat_high": {"ru": "🔴 Высокий приоритет", "kk": "🔴 Жоғары басымдық", "en": "🔴 High priority", "tr": "🔴 Yüksek öncelik"},
    "excel_title": {"ru": "ЖУРНАЛ ЖАЛОБ И ОБРАЩЕНИЙ", "kk": "ШАҒЫМДАР ЖУРНАЛЫ", "en": "COMPLAINTS LOG", "tr": "ŞİKAYET GÜNLÜĞÜ"},
    "excel_period": {"ru": "Период:", "kk": "Кезең:", "en": "Period:", "tr": "Dönem:"},
    "excel_h_date": {"ru": "Дата", "kk": "Күні", "en": "Date", "tr": "Tarih"},
    "excel_h_source": {"ru": "Источник", "kk": "Көзі", "en": "Source", "tr": "Kaynak"},
    "excel_h_note": {"ru": "Комментарий", "kk": "Түсініктеме", "en": "Comment", "tr": "Yorum"},
    "excel_h_cat": {"ru": "Категория", "kk": "Санат", "en": "Category", "tr": "Kategori"},
    "excel_h_prio": {"ru": "Приоритет", "kk": "Басымдық", "en": "Priority", "tr": "Öncelik"},
    "excel_h_status": {"ru": "Статус", "kk": "Мәртебе", "en": "Status", "tr": "Durum"},
    "excel_h_guest": {"ru": "Гость/№", "kk": "Қонақ/№", "en": "Guest/№", "tr": "Misafir/No"},
    "excel_h_dep": {"ru": "Отдел", "kk": "Бөлім", "en": "Dept", "tr": "Departman"},
    "excel_h_asn": {"ru": "Исполнитель", "kk": "Орындаушы", "en": "Assignee", "tr": "Sorumlu"},
    "excel_h_dl": {"ru": "Срок", "kk": "Мерзім", "en": "Deadline", "tr": "Son tarih"},
    "excel_h_desc": {"ru": "Описание", "kk": "Сипаттама", "en": "Description", "tr": "Açıklama"},
    "excel_h_meas": {"ru": "Меры", "kk": "Шаралар", "en": "Measures", "tr": "Önlemler"},
    "excel_h_res": {"ru": "Итог", "kk": "Нәтиже", "en": "Result", "tr": "Sonuç"},
    "excel_saved": {"ru": "Сохранено:", "kk": "Сақталды:", "en": "Saved:", "tr": "Kaydedildi:"},
    "excel_error": {"ru": "Ошибка", "kk": "Қате", "en": "Error", "tr": "Hata"},
    "excel_install": {"ru": "Установите: pip install openpyxl", "kk": "Орнатыңыз: pip install openpyxl", "en": "Install: pip install openpyxl", "tr": "Yükleyin: pip install openpyxl"},
}

def tt(key):
    d = T.get(key)
    return d.get(get_language(), d.get("ru", key)) if d else key

def db():
    return sqlite3.connect(DB_PATH)

def ensure_table():
    c = db()
    c.execute("""CREATE TABLE IF NOT EXISTS complaints (
        id INTEGER PRIMARY KEY AUTOINCREMENT, date TEXT, source TEXT, source_note TEXT DEFAULT '',
        category TEXT, priority TEXT, status TEXT, guest TEXT, department TEXT, assignee TEXT,
        deadline TEXT, description TEXT, measures TEXT, result TEXT,
        created_at TEXT DEFAULT (datetime('now','localtime'))
    )""")
    try:
        c.execute("ALTER TABLE complaints ADD COLUMN source_note TEXT DEFAULT ''")
    except sqlite3.OperationalError:
        pass
    c.commit()
    c.close()

def month_bounds():
    t = date.today()
    f = t.replace(day=1).isoformat()
    to = t.replace(month=t.month + 1, day=1) - timedelta(days=1) if t.month != 12 else t.replace(year=t.year + 1, month=1, day=1) - timedelta(days=1)
    return f, to.isoformat()

def build_complaints_page(master):
    ComplaintsPage(master).pack(fill="both", expand=True)


class ComplaintsPage(ctk.CTkFrame):
    def __init__(self, master, **kw):
        super().__init__(master, **kw)
        self.configure(fg_color="transparent")
        ensure_table()
        f, t = month_bounds()
        self.from_var = ctk.StringVar(value=f)
        self.to_var = ctk.StringVar(value=t)
        self.filter_cat = ctk.StringVar(value=tt("all"))
        self.filter_st = ctk.StringVar(value=tt("all"))
        self.rebuild()

    def rebuild(self):
        for w in self.winfo_children():
            w.destroy()

        head = ctk.CTkFrame(self)
        head.pack(fill="x", padx=10, pady=(10, 5))
        ctk.CTkLabel(head, text=tt("title"), font=ctk.CTkFont(size=20, weight="bold")).pack(side="left", padx=10, pady=8)
        ctk.CTkButton(head, text=tt("export"), width=110, fg_color="#1d4ed8", hover_color="#1e40af",
                      command=self.export_excel).pack(side="right", padx=10)
        ctk.CTkButton(head, text=tt("add"), width=180, fg_color="#0891b2", hover_color="#0e7490",
                      command=self.dialog).pack(side="right", padx=10)

        stats = self._stats()
        sf = ctk.CTkFrame(self, fg_color="transparent")
        sf.pack(fill="x", padx=10, pady=5)
        for color, label_key, val in [
            ("#1e40af", "stat_total", stats["total"]),
            ("#dc2626", "stat_new", stats["new"]),
            ("#ca8a04", "stat_work", stats["work"]),
            ("#16a34a", "stat_done", stats["done"]),
            ("#7c3aed", "stat_high", stats["high"]),
        ]:
            card = ctk.CTkFrame(sf, fg_color=color, corner_radius=10)
            card.pack(side="left", fill="x", expand=True, padx=3)
            ctk.CTkLabel(card, text=str(val), font=ctk.CTkFont(size=22, weight="bold")).pack(padx=10, pady=(6, 0))
            ctk.CTkLabel(card, text=tt(label_key), font=ctk.CTkFont(size=11)).pack(padx=10, pady=(0, 6))

        pf = ctk.CTkFrame(self)
        pf.pack(fill="x", padx=10, pady=5)
        ctk.CTkLabel(pf, text=tt("from")).pack(side="left", padx=(10, 4))
        ctk.CTkEntry(pf, textvariable=self.from_var, width=110).pack(side="left")
        ctk.CTkLabel(pf, text=tt("to")).pack(side="left", padx=(10, 4))
        ctk.CTkEntry(pf, textvariable=self.to_var, width=110).pack(side="left")
        ctk.CTkOptionMenu(pf, values=[tt("all")] + CATEGORIES, variable=self.filter_cat, width=180).pack(side="left", padx=10)
        ctk.CTkOptionMenu(pf, values=[tt("all")] + STATUSES, variable=self.filter_st, width=140).pack(side="left")
        ctk.CTkButton(pf, text="🔎", width=40, fg_color="gray25", command=self.rebuild).pack(side="left", padx=10)

        card = ctk.CTkFrame(self)
        card.pack(fill="both", expand=True, padx=10, pady=5)
        box = ctk.CTkScrollableFrame(card, height=360)
        box.pack(fill="both", expand=True, padx=10, pady=10)

        rows = self._fetch()
        if not rows:
            ctk.CTkLabel(box, text=tt("no_rec"), text_color="gray").pack(pady=10)
            return

        lang = get_language()

        for r in rows:
            rid, d, src, srcnote, cat, prio, st, guest, dep, asn, dl, desc, meas, res = r
            row = ctk.CTkFrame(box, fg_color="gray15")
            row.pack(fill="x", pady=3)

            # Переводим отображаемые значения
            cat_display = get_complaint_translation(cat, lang)
            prio_display = get_complaint_translation(prio, lang)
            st_display = get_complaint_translation(st, lang)
            src_display = get_complaint_translation(src, lang)

            prio_col = {"🔴 Высокая": "#dc2626", "🟡 Средняя": "#ca8a04"}.get(prio, "#16a34a")
            st_col = {"🆕 Новая": "#dc2626", "⚙️ В работе": "#ca8a04", "✅ Решена": "#16a34a"}.get(st, "#6b7280")

            top = ctk.CTkFrame(row, fg_color="transparent")
            top.pack(fill="x", padx=8, pady=(6, 0))
            ctk.CTkLabel(top, text=f"{d} • {cat_display}", font=ctk.CTkFont(size=13, weight="bold"),
                         anchor="w").pack(side="left")
            ctk.CTkLabel(top, text=prio_display, text_color=prio_col,
                         font=ctk.CTkFont(size=12, weight="bold")).pack(side="right", padx=(4, 0))
            ctk.CTkLabel(top, text=st_display, text_color=st_col,
                         font=ctk.CTkFont(size=12, weight="bold")).pack(side="right")

            src_txt = f"📥 {src_display}" + (f" • 💬 {srcnote}" if srcnote else "") + (f" •  {guest}" if guest else "")
            ctk.CTkLabel(row, text=src_txt, text_color="gray", anchor="w",
                         justify="left", wraplength=880).pack(anchor="w", padx=8)
            ctk.CTkLabel(row, text=f"🏢 {dep}" + (f" • 🔧 {asn}" if asn else "") + (f" • ⏰ {tt('to')} {dl}" if dl else ""),
                         text_color="#7fb6c9", anchor="w").pack(anchor="w", padx=8)
            ctk.CTkLabel(row, text=desc, anchor="w", justify="left", wraplength=880).pack(anchor="w", padx=8)
            if meas:
                ctk.CTkLabel(row, text=f"🛠 {meas}", text_color="#86efac", anchor="w",
                             justify="left", wraplength=880).pack(anchor="w", padx=8)
            if res:
                ctk.CTkLabel(row, text=f"✅ {res}", text_color="#93c5fd", anchor="w",
                             justify="left", wraplength=880).pack(anchor="w", padx=8, pady=(0, 6))

            ctk.CTkButton(row, text="️", width=36, fg_color="#7f1d1d", hover_color="#991b1b",
                          command=lambda i=rid: self.delete(i)).pack(side="right", padx=6, pady=6)
            ctk.CTkButton(row, text="✏️", width=36, fg_color="gray25",
                          command=lambda rec=r: self.dialog(rec)).pack(side="right", pady=6)

    def _where(self):
        conds = ["date BETWEEN ? AND ?"]
        args = [self.from_var.get(), self.to_var.get()]
        if self.filter_cat.get() != tt("all"):
            conds.append("category = ?")
            args.append(self.filter_cat.get())
        if self.filter_st.get() != tt("all"):
            conds.append("status = ?")
            args.append(self.filter_st.get())
        return " AND ".join(conds), args

    def _fetch(self):
        w, a = self._where()
        return db().execute(
            f"SELECT id, date, source, source_note, category, priority, status, guest, department, "
            f"assignee, deadline, description, measures, result "
            f"FROM complaints WHERE {w} ORDER BY date DESC, id DESC", a).fetchall()

    def _stats(self):
        rows = self._fetch()
        return {
            "total": len(rows),
            "new": sum(1 for r in rows if r[6] == "🆕 Новая"),
            "work": sum(1 for r in rows if r[6] == "⚙️ В работе"),
            "done": sum(1 for r in rows if r[6] == "✅ Решена"),
            "high": sum(1 for r in rows if r[5] == "🔴 Высокая"),
        }

    def delete(self, rid):
        if messagebox.askyesno(tt("warning"), tt("del")):
            c = db(); c.execute("DELETE FROM complaints WHERE id=?", (rid,)); c.commit(); c.close()
            self.rebuild()

    def dialog(self, rec=None):
        win = ctk.CTkToplevel(self)
        win.title(tt("title"))
        win.geometry("560x780")
        win.grab_set()

        lang = get_language()

        # Переведенные списки
        tr_sources = translate_list(SOURCES, lang)
        tr_categories = translate_list(CATEGORIES, lang)
        tr_priorities = translate_list(PRIORITIES, lang)
        tr_statuses = translate_list(STATUSES, lang)
        tr_types = translate_list(COMPLAINT_TYPES, lang)

        # Обратный маппинг: переведенное -> оригинал
        reverse_map = {}
        for orig_list, tr_list in [
            (SOURCES, tr_sources),
            (CATEGORIES, tr_categories),
            (PRIORITIES, tr_priorities),
            (STATUSES, tr_statuses),
            (COMPLAINT_TYPES, tr_types),
        ]:
            for orig, tr in zip(orig_list, tr_list):
                reverse_map[tr] = orig

        def find_translated(orig_value, orig_list, tr_list):
            try:
                idx = orig_list.index(orig_value)
                return tr_list[idx]
            except (ValueError, IndexError):
                return tr_list[0] if tr_list else ""

        g = rec or ("", date.today().isoformat(), SOURCES[0], "", CATEGORIES[0], PRIORITIES[1], STATUSES[0],
                    "", "", "", "", "", "", "")

        scroll = ctk.CTkScrollableFrame(win, corner_radius=10)
        scroll.pack(fill="both", expand=True, padx=12, pady=(12, 6))

        ctk.CTkLabel(scroll, text=tt("date")).pack(anchor="w", pady=(4, 0))
        e_date = ctk.CTkEntry(scroll); e_date.insert(0, g[1]); e_date.pack(fill="x")

        ctk.CTkLabel(scroll, text=tt("source")).pack(anchor="w", pady=(8, 0))
        m_src = ctk.CTkOptionMenu(scroll, values=tr_sources)
        m_src.set(find_translated(g[2], SOURCES, tr_sources))
        m_src.pack(fill="x")

        ctk.CTkLabel(scroll, text=tt("source_note")).pack(anchor="w", pady=(8, 0))
        e_srcnote = ctk.CTkEntry(scroll, placeholder_text="напр.: ссылка на отзыв, № комнаты, ФИО сообщившего")
        e_srcnote.insert(0, g[3]); e_srcnote.pack(fill="x")

        ctk.CTkLabel(scroll, text=tt("category")).pack(anchor="w", pady=(8, 0))
        m_cat = ctk.CTkOptionMenu(scroll, values=tr_categories)
        m_cat.set(find_translated(g[4], CATEGORIES, tr_categories))
        m_cat.pack(fill="x")

        r3 = ctk.CTkFrame(scroll, fg_color="transparent"); r3.pack(fill="x", pady=(8, 0))
        ctk.CTkLabel(r3, text=tt("priority")).pack(side="left", padx=(0, 6))
        m_prio = ctk.CTkOptionMenu(r3, values=tr_priorities, width=140)
        m_prio.set(find_translated(g[5], PRIORITIES, tr_priorities))
        m_prio.pack(side="left")
        ctk.CTkLabel(r3, text=tt("status")).pack(side="left", padx=(16, 6))
        m_st = ctk.CTkOptionMenu(r3, values=tr_statuses, width=140)
        m_st.set(find_translated(g[6], STATUSES, tr_statuses))
        m_st.pack(side="left")

        ctk.CTkLabel(scroll, text=tt("guest")).pack(anchor="w", pady=(8, 0))
        e_guest = ctk.CTkEntry(scroll); e_guest.insert(0, g[7]); e_guest.pack(fill="x")

        ctk.CTkLabel(scroll, text=tt("department")).pack(anchor="w", pady=(8, 0))
        e_dep = ctk.CTkEntry(scroll); e_dep.insert(0, g[8]); e_dep.pack(fill="x")

        ctk.CTkLabel(scroll, text=tt("assignee")).pack(anchor="w", pady=(8, 0))
        e_asn = ctk.CTkEntry(scroll); e_asn.insert(0, g[9]); e_asn.pack(fill="x")

        ctk.CTkLabel(scroll, text=tt("deadline")).pack(anchor="w", pady=(8, 0))
        e_dl = ctk.CTkEntry(scroll); e_dl.insert(0, g[10]); e_dl.pack(fill="x")

        ctk.CTkLabel(scroll, text=tt("description")).pack(anchor="w", pady=(8, 0))
        c_desc = ctk.CTkComboBox(scroll, values=tr_types)
        c_desc.set(find_translated(g[11], COMPLAINT_TYPES, tr_types) if g[11] else "")
        c_desc.pack(fill="x")

        ctk.CTkLabel(scroll, text=tt("measures")).pack(anchor="w", pady=(8, 0))
        t_meas = ctk.CTkTextbox(scroll, height=60); t_meas.pack(fill="x")
        t_meas.insert("1.0", g[12])

        ctk.CTkLabel(scroll, text=tt("result")).pack(anchor="w", pady=(8, 0))
        t_res = ctk.CTkTextbox(scroll, height=60); t_res.pack(fill="x")
        t_res.insert("1.0", g[13])

        def save():
            desc_trans = c_desc.get().strip()
            desc_orig = reverse_map.get(desc_trans, desc_trans)
            if not desc_orig:
                messagebox.showwarning(tt("warning"), tt("need_text")); return

            vals = (
                e_date.get(),
                reverse_map.get(m_src.get(), m_src.get()),
                e_srcnote.get().strip(),
                reverse_map.get(m_cat.get(), m_cat.get()),
                reverse_map.get(m_prio.get(), m_prio.get()),
                reverse_map.get(m_st.get(), m_st.get()),
                e_guest.get(), e_dep.get(), e_asn.get(), e_dl.get(),
                desc_orig,
                t_meas.get("1.0", "end").strip(),
                t_res.get("1.0", "end").strip()
            )
            c = db()
            if rec:
                c.execute("UPDATE complaints SET date=?, source=?, source_note=?, category=?, priority=?, status=?, "
                          "guest=?, department=?, assignee=?, deadline=?, description=?, measures=?, result=? WHERE id=?",
                          vals + (rec[0],))
            else:
                c.execute("INSERT INTO complaints (date, source, source_note, category, priority, status, guest, "
                          "department, assignee, deadline, description, measures, result) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                          vals)
            c.commit(); c.close()
            win.destroy(); self.rebuild()

        ctk.CTkButton(win, text=tt("save"), command=save, height=40,
                      fg_color="#0891b2", hover_color="#0e7490").pack(fill="x", padx=12, pady=(0, 12))

    def export_excel(self):
        rows = self._fetch()
        if not rows:
            messagebox.showwarning(tt("warning"), tt("no_rec")); return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            initialfile=f"complaints_{self.from_var.get()}_{self.to_var.get()}.xlsx")
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, Side, PatternFill
        except ImportError:
            messagebox.showerror(tt("excel_error"), tt("excel_install")); return

        wb = Workbook(); ws = wb.active; ws.title = tt("title")
        wrap = Alignment(wrap_text=True, vertical="top")
        thin = Border(*[Side(style="thin")] * 4)
        bold = Font(bold=True)
        yellow = PatternFill("solid", fgColor="FEC50C")

        ws.merge_cells("A1:M1")
        ws["A1"] = tt("excel_title")
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].fill = yellow
        ws["A2"] = f"{tt('excel_period')} {self.from_var.get()} — {self.to_var.get()}"

        headers = [tt("excel_h_date"), tt("excel_h_source"), tt("excel_h_note"), tt("excel_h_cat"),
                   tt("excel_h_prio"), tt("excel_h_status"), tt("excel_h_guest"), tt("excel_h_dep"),
                   tt("excel_h_asn"), tt("excel_h_dl"), tt("excel_h_desc"), tt("excel_h_meas"), tt("excel_h_res")]
        for j, h in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=j, value=h); cell.font = bold; cell.border = thin

        prio_fill = {"🔴 Высокая": "FCA5A5", "🟡 Средняя": "FDE68A", "🟢 Низкая": "BBF7D0"}
        st_fill = {"🆕 Новая": "FCA5A5", "⚙️ В работе": "FDE68A", "✅ Решена": "BBF7D0", "❌ Отклонена": "D1D5DB"}

        lang = get_language()
        for i, r in enumerate(rows, start=5):
            vals = list(r[1:14])
            # Переводим значения для Excel
            vals[1] = get_complaint_translation(vals[1], lang)  # source
            vals[3] = get_complaint_translation(vals[3], lang)  # category
            vals[4] = get_complaint_translation(vals[4], lang)  # priority
            vals[5] = get_complaint_translation(vals[5], lang)  # status

            for j, v in enumerate(vals, start=1):
                cell = ws.cell(row=i, column=j, value=v); cell.alignment = wrap; cell.border = thin

            p = vals[4]
            if p in prio_fill:
                ws.cell(row=i, column=5).fill = PatternFill("solid", fgColor=prio_fill[p])
            s = vals[5]
            if s in st_fill:
                ws.cell(row=i, column=6).fill = PatternFill("solid", fgColor=st_fill[s])

        for j, w in enumerate([12, 24, 32, 22, 14, 14, 22, 20, 20, 12, 40, 30, 30], start=1):
            ws.column_dimensions[chr(64 + j)].width = w
        wb.save(path)
        messagebox.showinfo("✅", f"{tt('excel_saved')}\n{path}")
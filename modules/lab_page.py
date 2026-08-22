# -*- coding: utf-8 -*-
"""Журнал лабораторных исследований + отчёт 02-ИРПК (приказ № 62, приложение 2)."""
import sqlite3
from datetime import date
from pathlib import Path

import customtkinter as ctk
from tkinter import filedialog, messagebox

from modules.translations import get_language

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DB_PATH = PROJECT_ROOT / "database" / "sanepi.db"

SAMPLE_TYPES = ["Вода", "Смывы", "Воздух", "Сырьё", "Готовая продукция", "Постельные принадлежности",
                "Почва/песок", "Дезрастворы", "Физические факторы (шум, освещенность, микроклимат)"]
LAB_TYPES = ["С привлечением лаборатории", "На базе производственной лаборатории"]
RESULTS = ["Соответствует", "Не соответствует"]

T = {
    "title": {"ru": "🧪 Лабораторные исследования", "kk": "🧪 Зертханалық зерттеулер", "en": "🧪 Lab tests", "tr": "🧪 Laboratuvar"},
    "add": {"ru": "➕ Добавить исследование", "kk": "➕ Зерттеу қосу", "en": "➕ Add test", "tr": "➕ Ekle"},
    "export": {"ru": "📊 02-ИРПК (Excel)", "kk": "📊 02-ИРПК (Excel)", "en": "📊 02-IRPK (Excel)", "tr": "📊 02-IRPK (Excel)"},
    "date": {"ru": "Дата (ГГГГ-ММ-ДД):", "kk": "Күні (ЖЖЖЖ-АА-КК):", "en": "Date:", "tr": "Tarih:"},
    "object": {"ru": "Объект / подразделение:", "kk": "Нысан / бөлімше:", "en": "Object:", "tr": "Nesne:"},
    "sample": {"ru": "Объект исследования (проба):", "kk": "Зерттеу нысаны:", "en": "Sample type:", "tr": "Numune:"},
    "indicators": {"ru": "Показатели:", "kk": "Көрсеткіштер:", "en": "Indicators:", "tr": "Göstergeler:"},
    "lab_type": {"ru": "Кто проводил:", "kk": "Кім жүргізді:", "en": "Lab type:", "tr": "Kim:"},
    "lab_name": {"ru": "Название лаборатории:", "kk": "Зертхана атауы:", "en": "Lab name:", "tr": "Lab adı:"},
    "samples": {"ru": "Число проб/замеров:", "kk": "Сынама саны:", "en": "Samples:", "tr": "Numune sayısı:"},
    "result": {"ru": "Результат:", "kk": "Нәтиже:", "en": "Result:", "tr": "Sonuç:"},
    "measures": {"ru": "Принятые меры (при несоответствии):", "kk": "Қабылданған шаралар:", "en": "Measures:", "tr": "Önlemler:"},
    "save": {"ru": "💾 Сохранить", "kk": "💾 Сақтау", "en": "💾 Save", "tr": "💾 Kaydet"},
    "warning": {"ru": "⚠️ Внимание", "kk": "⚠️ Назар аударыңыз", "en": "⚠️ Warning", "tr": "⚠️ Uyarı"},
    "need_object": {"ru": "Укажите объект и показатели!", "kk": "Нысан мен көрсеткіштерді көрсетіңіз!", "en": "Enter object and indicators!", "tr": "Nesne girin!"},
    "del": {"ru": "Удалить запись?", "kk": "Жазбаны жою керек пе?", "en": "Delete?", "tr": "Silinsin mi?"},
    "no_rec": {"ru": "Записей за период нет.", "kk": "Кезеңде жазбалар жоқ.", "en": "No records.", "tr": "Kayıt yok."},
    "from": {"ru": "с", "kk": "бастап", "en": "from", "tr": "başlangıç"},
    "to": {"ru": "по", "kk": "дейін", "en": "to", "tr": "bitiş"},
}


def tt(key):
    d = T.get(key)
    if not d:
        return key
    return d.get(get_language(), d.get("ru", key))


def db():
    return sqlite3.connect(DB_PATH)


def ensure_table():
    c = db()
    c.execute("""CREATE TABLE IF NOT EXISTS lab_records (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date TEXT,
        object_name TEXT,
        sample_type TEXT,
        indicators TEXT,
        lab_type TEXT,
        lab_name TEXT,
        samples_count INTEGER DEFAULT 1,
        result TEXT,
        measures TEXT,
        created_at TEXT DEFAULT (datetime('now','localtime'))
    )""")
    c.commit()
    c.close()


def half_year_bounds():
    t = date.today()
    if t.month <= 6:
        return f"{t.year}-01-01", f"{t.year}-06-30"
    return f"{t.year}-07-01", f"{t.year}-12-31"


def build_lab_page(master):
    LabPage(master).pack(fill="both", expand=True)


class LabPage(ctk.CTkFrame):
    def __init__(self, master, **kwargs):
        super().__init__(master, **kwargs)
        self.configure(fg_color="transparent")
        ensure_table()
        f, t = half_year_bounds()
        self.from_var = ctk.StringVar(value=f)
        self.to_var = ctk.StringVar(value=t)
        self.rebuild()

    def rebuild(self):
        for w in self.winfo_children():
            w.destroy()
        head = ctk.CTkFrame(self)
        head.pack(fill="x", padx=10, pady=(10, 5))
        ctk.CTkLabel(head, text=tt("title"), font=ctk.CTkFont(size=20, weight="bold")).pack(side="left", padx=10, pady=8)
        ctk.CTkButton(head, text=tt("export"), width=170, fg_color="#1d4ed8", hover_color="#1e40af",
                      command=self.export_report).pack(side="right", padx=10)
        ctk.CTkButton(head, text=tt("add"), width=200, fg_color="#0891b2", hover_color="#0e7490",
                      command=self.dialog).pack(side="right", padx=10)

        pf = ctk.CTkFrame(self)
        pf.pack(fill="x", padx=10, pady=5)
        ctk.CTkLabel(pf, text=tt("from")).pack(side="left", padx=(10, 4))
        ctk.CTkEntry(pf, textvariable=self.from_var, width=110).pack(side="left")
        ctk.CTkLabel(pf, text=tt("to")).pack(side="left", padx=(10, 4))
        ctk.CTkEntry(pf, textvariable=self.to_var, width=110).pack(side="left")
        ctk.CTkButton(pf, text="🔎", width=40, fg_color="gray25", command=self.rebuild).pack(side="left", padx=10)

        card = ctk.CTkFrame(self)
        card.pack(fill="both", expand=True, padx=10, pady=5)
        box = ctk.CTkScrollableFrame(card, height=320)
        box.pack(fill="both", expand=True, padx=10, pady=10)
        rows = self.get_records()
        if not rows:
            ctk.CTkLabel(box, text=tt("no_rec"), text_color="gray").pack(pady=10)
            return
        for r in rows:
            rid, d, obj, st, ind, lt, ln, cnt, res, meas = r
            row = ctk.CTkFrame(box, fg_color="gray15")
            row.pack(fill="x", pady=3)
            col = "#16A34A" if res == "Соответствует" else "#DC2626"
            ctk.CTkLabel(row, text=f"{d} • {obj} • {st}", font=ctk.CTkFont(size=13, weight="bold"),
                         anchor="w").pack(anchor="w", padx=8, pady=(6, 0))
            ctk.CTkLabel(row, text=f"{ind} • {cnt} проб • {lt}" + (f" ({ln})" if ln else ""),
                         text_color="gray", anchor="w", justify="left", wraplength=880).pack(anchor="w", padx=8)
            ctk.CTkLabel(row, text=res + (f" • меры: {meas}" if meas else ""), text_color=col,
                         font=ctk.CTkFont(size=12, weight="bold")).pack(anchor="w", padx=8, pady=(0, 6))
            ctk.CTkButton(row, text="🗑️", width=36, fg_color="#7f1d1d", hover_color="#991b1b",
                          command=lambda i=rid: self.delete(i)).pack(side="right", padx=6, pady=6)
            ctk.CTkButton(row, text="✏️", width=36, fg_color="gray25",
                          command=lambda rec=r: self.dialog(rec)).pack(side="right", pady=6)

    def get_records(self):
        return db().execute(
            "SELECT id, date, object_name, sample_type, indicators, lab_type, lab_name, samples_count, result, measures "
            "FROM lab_records WHERE date BETWEEN ? AND ? ORDER BY date",
            (self.from_var.get(), self.to_var.get())).fetchall()

    def delete(self, rid):
        if messagebox.askyesno(tt("warning"), tt("del")):
            c = db(); c.execute("DELETE FROM lab_records WHERE id=?", (rid,)); c.commit(); c.close()
            self.rebuild()

    def dialog(self, rec=None):
        win = ctk.CTkToplevel(self)
        win.title(tt("title"))
        win.geometry("470x640")
        win.grab_set()
        g = rec or ("", date.today().isoformat(), "", "", "", LAB_TYPES[0], "", 1, RESULTS[0], "")

        ctk.CTkLabel(win, text=tt("date")).pack(anchor="w", padx=20, pady=(10, 0))
        e_date = ctk.CTkEntry(win); e_date.insert(0, g[1]); e_date.pack(fill="x", padx=20)
        ctk.CTkLabel(win, text=tt("object")).pack(anchor="w", padx=20, pady=(8, 0))
        e_obj = ctk.CTkEntry(win); e_obj.insert(0, g[2]); e_obj.pack(fill="x", padx=20)
        ctk.CTkLabel(win, text=tt("sample")).pack(anchor="w", padx=20, pady=(8, 0))
        m_st = ctk.CTkOptionMenu(win, values=SAMPLE_TYPES); m_st.set(g[3] or SAMPLE_TYPES[0]); m_st.pack(fill="x", padx=20)
        ctk.CTkLabel(win, text=tt("indicators")).pack(anchor="w", padx=20, pady=(8, 0))
        e_ind = ctk.CTkEntry(win, placeholder_text="микробиология, сан-химия, БГКП…"); e_ind.insert(0, g[4]); e_ind.pack(fill="x", padx=20)
        ctk.CTkLabel(win, text=tt("lab_type")).pack(anchor="w", padx=20, pady=(8, 0))
        m_lt = ctk.CTkOptionMenu(win, values=LAB_TYPES); m_lt.set(g[5]); m_lt.pack(fill="x", padx=20)
        ctk.CTkLabel(win, text=tt("lab_name")).pack(anchor="w", padx=20, pady=(8, 0))
        e_ln = ctk.CTkEntry(win); e_ln.insert(0, g[6]); e_ln.pack(fill="x", padx=20)
        ctk.CTkLabel(win, text=tt("samples")).pack(anchor="w", padx=20, pady=(8, 0))
        e_cnt = ctk.CTkEntry(win); e_cnt.insert(0, str(g[7])); e_cnt.pack(fill="x", padx=20)
        ctk.CTkLabel(win, text=tt("result")).pack(anchor="w", padx=20, pady=(8, 0))
        m_res = ctk.CTkOptionMenu(win, values=RESULTS); m_res.set(g[8]); m_res.pack(fill="x", padx=20)
        ctk.CTkLabel(win, text=tt("measures")).pack(anchor="w", padx=20, pady=(8, 0))
        e_me = ctk.CTkEntry(win); e_me.insert(0, g[9]); e_me.pack(fill="x", padx=20)

        def save():
            if not e_obj.get().strip() or not e_ind.get().strip():
                messagebox.showwarning(tt("warning"), tt("need_object")); return
            try:
                cnt = int(e_cnt.get() or 1)
            except ValueError:
                cnt = 1
            c = db()
            if rec:
                c.execute("UPDATE lab_records SET date=?, object_name=?, sample_type=?, indicators=?, lab_type=?, lab_name=?, samples_count=?, result=?, measures=? WHERE id=?",
                          (e_date.get(), e_obj.get(), m_st.get(), e_ind.get(), m_lt.get(), e_ln.get(), cnt, m_res.get(), e_me.get(), rec[0]))
            else:
                c.execute("INSERT INTO lab_records (date, object_name, sample_type, indicators, lab_type, lab_name, samples_count, result, measures) VALUES (?,?,?,?,?,?,?,?,?)",
                          (e_date.get(), e_obj.get(), m_st.get(), e_ind.get(), m_lt.get(), e_ln.get(), cnt, m_res.get(), e_me.get()))
            c.commit(); c.close()
            win.destroy(); self.rebuild()
        ctk.CTkButton(win, text=tt("save"), command=save).pack(pady=12)

    # ---------- 02-ИРПК ----------
    def export_report(self):
        recs = self.get_records()
        if not recs:
            messagebox.showwarning(tt("warning"), tt("no_rec")); return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            initialfile=f"02-ИРПК_{self.from_var.get()}_{self.to_var.get()}.xlsx")
        if not path:
            return
        try:
            build_02_irpk(recs, self.from_var.get(), self.to_var.get(), path)
        except ImportError:
            messagebox.showerror("Ошибка", "Нужна библиотека openpyxl:\npip install openpyxl")
            return
        messagebox.showinfo("✅", f"Отчёт 02-ИРПК сохранён:\n{path}")


def build_02_irpk(recs, d_from, d_to, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = Workbook(); ws = wb.active; ws.title = "02-ИРПК"
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(*[Side(style="thin")] * 4)
    bold = Font(bold=True)

    ws.merge_cells("A1:F1"); ws["A1"] = "Информация о результатах производственного контроля (форма 02-ИРПК)"
    ws["A1"].font = bold; ws["A1"].alignment = wrap
    ws.merge_cells("A2:F2"); ws["A2"] = f"Отчетный период: с {d_from} по {d_to} (полугодие)"
    ws["A2"].alignment = wrap
    ws["A3"] = "Наименование: ____________________________________  ИИН/БСН: ____________"
    ws.merge_cells("A4:F4"); ws["A4"] = "Адрес: ____________________ Телефон: ____________ E-mail: ____________________"

    headers = ["№ п/п",
               "Сведения о лице, осуществляющем производственный контроль, на базе производственной лаборатории объекта",
               "Сведения о лице, осуществляющем производственный контроль, с привлечением лаборатории (испытательного центра)",
               "Всего исследовано (перечислить объекты внешней среды и число проб – сырье, готовая продукция, смывы, воздух и другие)",
               "Выявлено несоответствий (перечислить показатели безопасности, по которым выявлено несоответствие – БГКП, патогенная флора, токсические вещества и другие)",
               "Принятые меры и проведенные мероприятия по устранению"]
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=6, column=j, value=h); cell.font = bold; cell.alignment = wrap; cell.border = thin

    groups = {}
    for rid, d, obj, st, ind, lt, ln, cnt, res, meas in recs:
        g = groups.setdefault(st, {"own": "", "ext": "", "count": 0, "objs": set(), "bad": [], "meas": []})
        if lt == LAB_TYPES[1]:
            g["own"] = ln or "имеется"
        else:
            g["ext"] = ln or g["ext"] or "привлеченная лаборатория"
        g["count"] += cnt or 0
        g["objs"].add(obj)
        if res == "Не соответствует":
            g["bad"].append(f"{ind} ({obj})")
            if meas:
                g["meas"].append(meas)

    row_i = 7
    for n, (st, g) in enumerate(groups.items(), start=1):
        vals = [n, g["own"] or "—", g["ext"] or "—",
                f"{st}: {g['count']} проб/замеров ({'; '.join(sorted(g['objs']))})",
                "; ".join(g["bad"]) or "не выявлено",
                "; ".join(g["meas"]) or "—"]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=row_i, column=j, value=v); cell.alignment = wrap; cell.border = thin
        row_i += 1

    widths = [6, 28, 28, 40, 35, 35]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.merge_cells("A1:F1"); ws.merge_cells("A2:F2"); ws.merge_cells("A4:F4")
    wb.save(path)